import os
import math
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import streamlit as st

# Importações para o ReportLab (PDF)
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Image as RLImage
from reportlab.lib import colors

def set_cell_margins(cell, top=20, bottom=20, left=20, right=20):
    """Define margens internas reduzidas na célula do Word para otimizar espaço."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def gerar_word_personalizado(image_path, largura_tag, altura_tag, qtd_tags):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Cm(0.4)
        section.bottom_margin = Cm(0.4)
        section.left_margin = Cm(0.4)
        section.right_margin = Cm(0.4)

    largura_pagina_util = 20.0
    cols = max(1, int(largura_pagina_util // largura_tag))
    rows = math.ceil(qtd_tags / cols)

    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    tag_contador = 0
    for i in range(rows):
        row = table.rows[i]
        row.height = Cm(altura_tag)
        row.height_rule = 1
        for j in range(cols):
            if tag_contador >= qtd_tags:
                break
            cell = row.cells[j]
            cell.width = Cm(largura_tag)
            set_cell_margins(cell)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run()
            run.add_picture(image_path, width=Cm(largura_tag * 0.95))
            tag_contador += 1

    output_path = "tags_personalizadas.docx"
    doc.save(output_path)
    return output_path

def gerar_pdf_personalizado(image_path, largura_tag, altura_tag, qtd_tags):
    output_path = "tags_personalizadas.pdf"
    
    # Margens e conversão de cm para pontos (1 cm = 28.3465 pontos)
    pt_largura = largura_tag * 28.3465
    pt_altura = altura_tag * 28.3465
    
    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=15,
        rightMargin=15,
        topMargin=15,
        bottomMargin=15
    )
    
    # Largura útil da página A4 (~565 pontos)
    largura_util_pt = 565
    cols = max(1, int(largura_util_pt // pt_largura))
    rows = math.ceil(qtd_tags / cols)
    
    data = []
    tag_contador = 0
    for r in range(rows):
        row_cells = []
        for c in range(cols):
            if tag_contador < qtd_tags:
                img = RLImage(image_path, width=pt_largura * 0.92, height=pt_altura * 0.92)
                row_cells.append(img)
                tag_contador += 1
            else:
                row_cells.append("")
        data.append(row_cells)
        
    table = Table(data, colWidths=[pt_largura] * cols, rowHeights=[pt_altura] * rows)
    table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.transparent),
        ('BOX', (0,0), (-1,-1), 0.5, colors.transparent),
    ]))
    
    doc.build([table])
    return output_path

# Interface Streamlit
st.title("🧁 Gerador Personalizado de Tags")
st.write("Personalize o tamanho e a quantidade de tags à sua medida para impressão em Word ou PDF.")

st.sidebar.header("⚙️ Definições de Impressão")
largura_tag = st.sidebar.slider("Largura da Tag (cm)", min_value=2.0, max_value=10.0, value=4.0, step=0.5)
altura_tag = st.sidebar.slider("Altura da Tag (cm)", min_value=2.0, max_value=10.0, value=4.0, step=0.5)
qtd_tags = st.sidebar.number_input("Quantidade Total de Tags", min_value=1, max_value=100, value=35, step=1)

uploaded_file = st.file_uploader("Escolha a imagem da tag (PNG ou JPG)", type=["png", "jpg", "jpeg"], key="tag_uploader_custom_both")

if uploaded_file is not None:
    temp_image_path = "temp_tag.png"
    with open(temp_image_path, "wb") as f:
        f.write(uploaded_file.getbuffer())
    
    st.success("Imagem carregada com sucesso!")
    st.info(f"Configuração selecionada: **{qtd_tags} tags** com tamanho **{largura_tag}x{altura_tag} cm**.")
    
    col1, col2 = st.columns(2)
    
    with col1:
        if st.button("Gerar Documento Word (.docx)"):
            with st.spinner("A criar documento Word..."):
                docx_file = gerar_word_personalizado(temp_image_path, largura_tag, altura_tag, int(qtd_tags))
                with open(docx_file, "rb") as f:
                    st.download_button(
                        label="📥 Descarregar Word",
                        data=f,
                        file_name=f"Tags_{int(largura_tag)}x{int(altura_tag)}cm_{qtd_tags}un.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key="download_docx"
                    )
                    
    with col2:
        if st.button("Gerar Documento PDF (.pdf)"):
            with st.spinner("A criar documento PDF..."):
                pdf_file = gerar_pdf_personalizado(temp_image_path, largura_tag, altura_tag, int(qtd_tags))
                with open(pdf_file, "rb") as f:
                    st.download_button(
                        label="📥 Descarregar PDF",
                        data=f,
                        file_name=f"Tags_{int(largura_tag)}x{int(altura_tag)}cm_{qtd_tags}un.pdf",
                        mime="application/pdf",
                        key="download_pdf"
                    )
